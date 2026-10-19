from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from pdf2excel import extract_tables, to_csv, to_xlsx
from pdf2excel.cli import main

S = Path(__file__).resolve().parent.parent / "samples"


def test_xlsx_sheets_and_numbers(tmp_path):
    out = tmp_path / "inv.xlsx"
    to_xlsx(extract_tables(str(S / "sample_invoice.pdf")), out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["p1_t1", "p1_t2"]
    ws = wb["p1_t1"]
    assert [c.value for c in ws[1]] == ["Item", "Description", "Qty", "Unit Price", "Amount"]
    assert ws["E2"].value == 125.0  # converted to number


def test_xlsx_thousands_separator(tmp_path):
    out = tmp_path / "bank.xlsx"
    to_xlsx(extract_tables(str(S / "sample_bank_statement.pdf")), out)
    df = pd.read_excel(out, sheet_name="p2_t1")
    assert pd.api.types.is_numeric_dtype(df["Balance"])
    assert df["Balance"].iloc[0] == pytest.approx(-1437.42)


def test_csv_single_and_multi(tmp_path):
    p = to_csv(extract_tables(str(S / "sample_sales_report.pdf")), tmp_path / "r.csv")
    assert p == [tmp_path / "r.csv"]
    assert list(pd.read_csv(p[0]).columns) == ["Region", "Q1", "Q2", "Q3", "Q4"]
    ps = to_csv(extract_tables(str(S / "sample_bank_statement.pdf")), tmp_path / "b.csv")
    assert [x.name for x in ps] == ["b_p1_t1.csv", "b_p2_t1.csv"]


def test_empty_raises(tmp_path):
    with pytest.raises(ValueError):
        to_xlsx([], tmp_path / "x.xlsx")
    with pytest.raises(ValueError):
        to_csv([], tmp_path / "x.csv")


def test_cli_xlsx(tmp_path, capsys):
    out = tmp_path / "o.xlsx"
    assert main([str(S / "sample_invoice.pdf"), "-o", str(out)]) == 0
    assert out.exists() and "2 sheet(s)" in capsys.readouterr().out


def test_cli_csv_pages(tmp_path):
    out = tmp_path / "o.csv"
    assert main([str(S / "sample_bank_statement.pdf"), "--csv", "-p", "1", "-o", str(out)]) == 0
    assert len(pd.read_csv(out)) == 25


def test_cli_errors(tmp_path):
    assert main([str(tmp_path / "missing.pdf")]) == 2
    assert main([str(S / "sample_invoice.pdf"), "-p", "5-2", "-o", str(tmp_path / "x.xlsx")]) == 2
    bad = tmp_path / "bad.pdf"
    bad.write_text("not a pdf")
    assert main([str(bad), "-o", str(tmp_path / "y.xlsx")]) == 1


def test_cli_merge(tmp_path, capsys):
    out = tmp_path / "m.xlsx"
    assert main([str(S / "sample_bank_statement.pdf"), "--merge", "-o", str(out)]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["p1_t1"]
    assert wb["p1_t1"].max_row == 41  # header + 40 rows
