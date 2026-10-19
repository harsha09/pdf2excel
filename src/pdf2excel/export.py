"""Export extracted tables to XLSX / CSV."""
from __future__ import annotations

import io
from pathlib import Path
from typing import List, Union

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .extract import ExtractedTable


def _numeric(df: pd.DataFrame) -> pd.DataFrame:
    """Convert columns that look numeric (e.g. '1,234.50') to numbers."""
    out = df.copy()
    for c in out.columns:
        s = out[c].astype(str).str.replace(",", "", regex=False).str.strip()
        nonblank = s[s != ""]
        if len(nonblank) == 0:
            continue
        conv = pd.to_numeric(nonblank, errors="coerce")
        if conv.notna().all():
            out[c] = pd.to_numeric(s.replace("", None), errors="coerce")
    return out


def to_xlsx(tables: List[ExtractedTable], dest: Union[str, Path, io.BytesIO], numeric: bool = True) -> None:
    """Write one sheet per table."""
    if not tables:
        raise ValueError("No tables to export")
    with pd.ExcelWriter(dest, engine="openpyxl") as xw:
        for t in tables:
            df = _numeric(t.df) if numeric else t.df
            df.to_excel(xw, sheet_name=t.name[:31], index=False)
            ws = xw.sheets[t.name[:31]]
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="DDEBF7")
            for i, col in enumerate(df.columns, 1):
                width = max([len(str(col))] + [len(str(v)) for v in df[col].tolist()])
                ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
            ws.freeze_panes = "A2"


def to_csv(tables: List[ExtractedTable], dest: Union[str, Path]) -> List[Path]:
    """Write CSV(s). One table -> dest itself; several -> dest_stem_p{page}_t{n}.csv."""
    if not tables:
        raise ValueError("No tables to export")
    dest = Path(dest)
    if len(tables) == 1:
        tables[0].df.to_csv(dest, index=False)
        return [dest]
    paths = []
    for t in tables:
        p = dest.with_name(f"{dest.stem}_{t.name}{dest.suffix or '.csv'}")
        t.df.to_csv(p, index=False)
        paths.append(p)
    return paths
